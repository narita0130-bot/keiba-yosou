#!/usr/bin/env python3
"""Logic@競馬 の注目馬（ユーザー提供の無料note分）を Excel / CSV にまとめて出す。

入力は data/marks_record.csv の source=Logic 行と data/baselines.json。
有料の「厳選」レースは提供されていないので含まない。

  python3 scripts/export_logic.py --out data/exports/logic_picks   （openpyxl が必要）
  → logic_picks.xlsx（概要・全予想・日別・人気別・競馬場別・指名数別）と logic_picks.csv
"""
import argparse
import csv
import datetime as dt
import json
import os
from collections import Counter, defaultdict

from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter

WEEK = "月火水木金土日"
HEAD_FILL = PatternFill("solid", fgColor="DDE5F0")


def band_of(ninki, bands):
    for lo, hi, name in bands:
        if lo <= ninki <= hi:
            return name
    return ""


def load(path):
    rows = []
    for r in csv.DictReader(open(path, encoding="utf-8-sig")):
        if r["source"] != "Logic":
            continue
        scratched = r["scratched"] == "1"
        chaku = int(r["chaku"]) if r["chaku"].isdigit() else None
        ninki = int(r["ninki"]) if r["ninki"].isdigit() else None
        odds = float(r["odds"]) if r["odds"] else None
        tan = int(round(odds * 100)) if chaku == 1 and odds else 0
        fuku = int(float(r["fuku"])) if r["fuku"] else 0
        d = dt.date.fromisoformat(r["date"])
        rows.append({
            "date": r["date"], "dow": WEEK[d.weekday()], "venue": r["venue"],
            "race_no": int(r["race_no"]), "race_id": r["race_id"],
            "umaban": int(r["umaban"]), "horse": r["horse"],
            "ninki": ninki, "odds": odds, "chaku": chaku,
            "field": int(r["field_size"]) if r["field_size"] else None,
            "scratched": scratched, "tan": tan, "fuku": fuku,
        })
    rows.sort(key=lambda x: (x["date"], x["venue"], x["race_no"], x["umaban"]))
    return rows


def summarize(rows, base, bands):
    """返還（取消）を除いた成績。無作為は同じ人気構成での期待回収率。"""
    live = [r for r in rows if not r["scratched"]]
    n = len(live)
    if n == 0:
        return None
    win = sum(r["chaku"] == 1 for r in live)
    top3 = sum(r["chaku"] is not None and r["chaku"] <= 3 for r in live)
    tan = sum(r["tan"] for r in live)
    fuku = sum(r["fuku"] for r in live)
    bt = bf = 0.0
    m = 0
    for r in live:
        if r["ninki"] is None:
            continue
        b = band_of(r["ninki"], bands)
        bt += base["tansho"]["by_band"][b]
        bf += base["fukusho"]["by_band"][b]
        m += 1
    return {
        "n": n, "win": win, "top3": top3,
        "win_rate": win / n, "top3_rate": top3 / n,
        "tan_ret": tan / (n * 100), "fuku_ret": fuku / (n * 100),
        "tan_pay": tan, "fuku_pay": fuku,
        "base_tan": bt / m / 100 if m else None,
        "base_fuku": bf / m / 100 if m else None,
        "scratched": len(rows) - n,
    }


STAT_HEAD = ["頭数", "1着", "3着内", "勝率", "3着内率",
             "単勝 払戻計", "単勝 回収率", "同人気の無作為(単)",
             "複勝 払戻計", "複勝 回収率", "同人気の無作為(複)"]


def stat_cells(s):
    return [s["n"], s["win"], s["top3"], s["win_rate"], s["top3_rate"],
            s["tan_pay"], s["tan_ret"], s["base_tan"],
            s["fuku_pay"], s["fuku_ret"], s["base_fuku"]]


PCT_COLS_IN_STAT = {3, 4, 6, 7, 9, 10}      # STAT_HEAD 内の百分率列（0始まり）
YEN_COLS_IN_STAT = {5, 8}


def write_table(ws, head, data, pct=(), yen=(), widths=None, start_row=1):
    for j, h in enumerate(head, 1):
        c = ws.cell(row=start_row, column=j, value=h)
        c.font = Font(bold=True)
        c.fill = HEAD_FILL
        c.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
    for i, row in enumerate(data, start_row + 1):
        for j, v in enumerate(row, 1):
            c = ws.cell(row=i, column=j, value=v)
            if (j - 1) in pct and isinstance(v, (int, float)):
                c.number_format = "0.0%"
            elif (j - 1) in yen and isinstance(v, (int, float)):
                c.number_format = "#,##0"
    ws.freeze_panes = ws.cell(row=start_row + 1, column=1)
    for j in range(1, len(head) + 1):
        w = (widths or {}).get(j - 1, 11)
        ws.column_dimensions[get_column_letter(j)].width = w


def grouped(rows, key, base, bands, order=None):
    g = defaultdict(list)
    for r in rows:
        g[key(r)].append(r)
    keys = order if order else sorted(g)
    out = []
    for k in keys:
        if k in g:
            s = summarize(g[k], base, bands)
            if s:
                out.append((k, s))
    return out


def main():
    ap = argparse.ArgumentParser()
    ap.add_argument("--marks", default="data/marks_record.csv")
    ap.add_argument("--base", default="data/baselines.json")
    ap.add_argument("--out", default="data/exports/logic_picks")
    args = ap.parse_args()

    base = json.load(open(args.base))
    bands = base["bands"]
    rows = load(args.marks)
    os.makedirs(os.path.dirname(args.out), exist_ok=True)

    # ---- CSV（1行1頭） ----
    cols = ["日付", "曜日", "競馬場", "R", "race_id", "馬番", "馬名", "人気", "人気帯",
            "単勝オッズ", "着順", "単勝払戻", "複勝払戻", "頭数", "備考"]

    def flat(r):
        note = "取消（返還）" if r["scratched"] else ("競走中止・除外等" if r["chaku"] is None else "")
        return [r["date"], r["dow"], r["venue"], r["race_no"], r["race_id"], r["umaban"],
                r["horse"], r["ninki"], band_of(r["ninki"], bands) if r["ninki"] else "",
                r["odds"], r["chaku"], r["tan"], r["fuku"], r["field"], note]

    with open(args.out + ".csv", "w", encoding="utf-8-sig", newline="") as f:
        w = csv.writer(f)
        w.writerow(cols)
        for r in rows:
            w.writerow(flat(r))

    wb = Workbook()

    # ---- 概要 ----
    ws = wb.active
    ws.title = "概要"
    s = summarize(rows, base, bands)
    days = sorted({r["date"] for r in rows})
    races = {r["race_id"] for r in rows}
    ws["A1"] = "Logic@競馬 注目馬（無料note）の成績まとめ"
    ws["A1"].font = Font(bold=True, size=14)
    info = [
        ("期間", f"{days[0]} 〜 {days[-1]}"),
        ("開催日数", len(days)),
        ("レース数", len(races)),
        ("指名頭数（取消を除く）", s["n"]),
        ("取消（返還）", s["scratched"]),
        ("出典", "ユーザー提供のnote（PDF）。有料の【厳選】レースは含まない"),
        ("払戻", "1頭100円で買った場合。単勝払戻は確定単勝オッズ×100、複勝は結果ページの払戻"),
        ("同人気の無作為", "同じ人気の馬を無作為に買った場合の期待回収率（data/baselines.json）。"
                       "これを上回って初めて「オッズより上手い」と言える"),
    ]
    for i, (k, v) in enumerate(info, 3):
        ws.cell(row=i, column=1, value=k).font = Font(bold=True)
        ws.cell(row=i, column=2, value=v)
    r0 = len(info) + 4
    write_table(ws, ["区分"] + STAT_HEAD, [["全体"] + stat_cells(s)],
                pct={j + 1 for j in PCT_COLS_IN_STAT}, yen={j + 1 for j in YEN_COLS_IN_STAT},
                widths={0: 22, 1: 14}, start_row=r0)
    ws.freeze_panes = None
    ws.column_dimensions["B"].width = 60

    # ---- 全予想 ----
    ws = wb.create_sheet("全予想")
    write_table(ws, cols, [flat(r) for r in rows], yen={11, 12},
                widths={0: 11, 4: 14, 6: 20, 14: 16})
    ws.auto_filter.ref = ws.dimensions

    pct = {j + 1 for j in PCT_COLS_IN_STAT}
    yen = {j + 1 for j in YEN_COLS_IN_STAT}

    # ---- 日別 ----
    ws = wb.create_sheet("日別")
    data = [[k, rows_dow] + stat_cells(v)
            for (k, v), rows_dow in
            ((kv, WEEK[dt.date.fromisoformat(kv[0]).weekday()])
             for kv in grouped(rows, lambda r: r["date"], base, bands))]
    write_table(ws, ["日付", "曜日"] + STAT_HEAD, data,
                pct={j + 2 for j in PCT_COLS_IN_STAT}, yen={j + 2 for j in YEN_COLS_IN_STAT},
                widths={0: 11, 1: 6})

    # ---- 人気別 ----
    ws = wb.create_sheet("人気別")
    order = [b[2] for b in bands]
    data = [[k] + stat_cells(v) for k, v in
            grouped([r for r in rows if r["ninki"]], lambda r: band_of(r["ninki"], bands),
                    base, bands, order)]
    write_table(ws, ["人気帯"] + STAT_HEAD, data, pct=pct, yen=yen, widths={0: 12})

    # ---- 競馬場別 ----
    ws = wb.create_sheet("競馬場別")
    data = [[k] + stat_cells(v) for k, v in grouped(rows, lambda r: r["venue"], base, bands)]
    data.sort(key=lambda x: -x[1])
    write_table(ws, ["競馬場"] + STAT_HEAD, data, pct=pct, yen=yen, widths={0: 10})

    # ---- 指名数別（1レースに何頭挙げたか） ----
    ws = wb.create_sheet("指名数別")
    per_race = Counter(r["race_id"] for r in rows)
    data = [[f"{k}頭指名"] + stat_cells(v) for k, v in
            grouped(rows, lambda r: per_race[r["race_id"]], base, bands)]
    write_table(ws, ["1レースの指名数"] + STAT_HEAD, data, pct=pct, yen=yen, widths={0: 14})

    wb.save(args.out + ".xlsx")
    print(f"wrote {args.out}.xlsx / {args.out}.csv  {len(days)}日 {len(races)}レース {len(rows)}頭")
    print(f"  勝率 {s['win_rate']:.1%}  3着内 {s['top3_rate']:.1%}  "
          f"単勝 {s['tan_ret']:.1%}（無作為 {s['base_tan']:.1%}）  "
          f"複勝 {s['fuku_ret']:.1%}（無作為 {s['base_fuku']:.1%}）")


if __name__ == "__main__":
    main()
